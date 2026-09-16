# # Final answer last 


# from flask import Flask, jsonify, request
# from flask_cors import CORS
# from pymongo import MongoClient
# from werkzeug.security import generate_password_hash, check_password_hash
# from werkzeug.utils import secure_filename
# from datetime import datetime
# import os
# import PyPDF2
# import docx
# import json
# import re 
# from groq import Groq
# from google import genai
# import pandas as pd
# from rec_courses import recommend_course
# from dotenv import load_dotenv

# app = Flask(__name__)
# CORS(app)
# load_dotenv()

# # Use environment variables with correct names
# connec_string = os.getenv('connec_string')
# groq_api = os.getenv('GROQ_API_KEY')
# gemini_api = os.getenv('GEMINI_API_KEY')

# print(f"Connection string exists: {bool(connec_string)}")
# print(f"Groq API key exists: {bool(groq_api)}")

# # Configure upload folder
# UPLOAD_FOLDER = 'uploads'
# app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# # Create uploads directory if it doesn't exist
# if not os.path.exists(UPLOAD_FOLDER):
#     os.makedirs(UPLOAD_FOLDER)
#     print(f"Created uploads directory: {UPLOAD_FOLDER}")

# # Initialize MongoDB
# client = None
# users_collection = None
# if connec_string:
#     try:
#         client = MongoClient(connec_string)
#         client.admin.command('ping')
#         print("MongoDB is connected")
#         db = client['UserTest']
#         users_collection = db['users']
#         print(f"Database connected: {db.name}")
#     except Exception as e:
#         print(f"MongoDB connection failed: {e}")
#         client = None
#         users_collection = None
# else:
#     print("No MongoDB connection string found")
#     users_collection = None

# # Initialize Groq client
# groq_client = None
# if groq_api:
#     try:
#         groq_client = Groq(api_key=groq_api)
#         print("Groq client initialized successfully")
#     except Exception as e:
#         print(f"Groq client initialization failed: {e}")
#         groq_client = None
# else:
#     print("No Groq API key found")
#     groq_client = None

# # ==================== AUTHENTICATION ROUTES ====================

# @app.route('/api/signup', methods=['POST'])
# def signup():
#     print("=== SIGNUP ENDPOINT CALLED ===")
#     try:
#         if not request.is_json:
#             return jsonify({'message': 'Request must be JSON'}), 400
        
#         data = request.get_json()
#         print(f"Received signup data: {data}")
        
#         name = data.get('name')
#         email = data.get('email')
#         password = data.get('password')
        
#         if not name or not email or not password:
#             return jsonify({'message': 'Name, email, and password are required'}), 400
        
#         # Check if user already exists
#         if users_collection is not None:
#             existing_user = users_collection.find_one({'email': email})
#             if existing_user:
#                 return jsonify({'message': 'User already exists with this email'}), 400
            
#             # Hash password and create user
#             hashed_password = generate_password_hash(password)
#             user_data = {
#                 'name': name,
#                 'email': email,
#                 'password': hashed_password,
#                 'created_at': datetime.now().isoformat(),
#                 'skills': [],
#                 'job': '',
#                 'tagline': 'A catchy tagline!'
#             }
            
#             result = users_collection.insert_one(user_data)
#             print(f"User created with ID: {result.inserted_id}")
            
#             return jsonify({
#                 'message': 'Signup successful',
#                 'user': {
#                     'name': name,
#                     'email': email
#                 }
#             }), 201
#         else:
#             return jsonify({'message': 'Database not available'}), 500
            
#     except Exception as e:
#         print(f"Signup error: {e}")
#         return jsonify({'message': 'Error creating user'}), 500

# @app.route('/api/login', methods=['POST'])
# def login():
#     print("=== LOGIN ENDPOINT CALLED ===")
#     try:
#         if not request.is_json:
#             return jsonify({'message': 'Request must be JSON'}), 400
        
#         data = request.get_json()
#         print(f"Received login data: {data}")
        
#         email = data.get('email')
#         password = data.get('password')
        
#         if not email or not password:
#             return jsonify({'message': 'Email and password are required'}), 400
        
#         if users_collection is not None:
#             user = users_collection.find_one({'email': email})
#             if user and check_password_hash(user['password'], password):
#                 user_data = {
#                     'name': user['name'],
#                     'email': user['email'],
#                     'id': str(user['_id'])
#                 }
#                 return jsonify({
#                     'message': 'Login successful',
#                     'user': user_data
#                 }), 200
#             else:
#                 return jsonify({'message': 'Invalid email or password'}), 401
#         else:
#             return jsonify({'message': 'Database not available'}), 500
            
#     except Exception as e:
#         print(f"Login error: {e}")
#         return jsonify({'message': 'Error during login'}), 500

# @app.route('/api/user/profile', methods=['GET'])
# def get_user_profile():
#     try:
#         email = request.args.get('email')
#         if not email:
#             return jsonify({'message': 'Email is required'}), 400
        
#         if users_collection is not None:
#             user = users_collection.find_one({'email': email})
#             if user:
#                 user_data = {
#                     'name': user.get('name', ''),
#                     'email': user.get('email', ''),
#                     'skills': user.get('skills', []),
#                     'job': user.get('job', ''),
#                     'tagline': user.get('tagline', 'A catchy tagline!')
#                 }
#                 return jsonify({'user': user_data}), 200
#             else:
#                 return jsonify({'message': 'User not found'}), 404
#         else:
#             return jsonify({'message': 'Database not available'}), 500
            
#     except Exception as e:
#         print(f"Profile error: {e}")
#         return jsonify({'message': 'Error fetching profile'}), 500

# @app.route('/api/user/profile', methods=['PUT'])
# def update_user_profile():
#     try:
#         data = request.get_json()
#         email = data.get('email')
#         updates = data.get('updates', {})
        
#         if not email or not updates:
#             return jsonify({'message': 'Email and updates are required'}), 400
        
#         if users_collection is not None:
#             # Remove password from updates if present
#             updates.pop('password', None)
            
#             result = users_collection.update_one(
#                 {'email': email},
#                 {'$set': updates}
#             )
            
#             if result.modified_count:
#                 return jsonify({'message': 'Profile updated successfully'}), 200
#             else:
#                 return jsonify({'message': 'No changes made or user not found'}), 400
#         else:
#             return jsonify({'message': 'Database not available'}), 500
            
#     except Exception as e:
#         print(f"Update profile error: {e}")
#         return jsonify({'message': 'Error updating profile'}), 500

# # ==================== SKILL ANALYZER ROUTE ====================

# @app.route('/api/skill-analyzer', methods=['POST'])
# def skill_analyzer():
#     print("=== SKILL ANALYZER ENDPOINT CALLED ===")
    
#     # Check if files are present
#     if 'resume' not in request.files:
#         print("Missing resume file")
#         return jsonify({'error': 'Resume file is required'}), 400
    
#     if 'job_description' not in request.files:
#         print("Missing job description file")
#         return jsonify({'error': 'Job description file is required'}), 400
    
#     resume_file = request.files['resume']
#     job_description_file = request.files['job_description']
    
#     # Check if files have names
#     if resume_file.filename == '':
#         return jsonify({'error': 'No resume file selected'}), 400
#     if job_description_file.filename == '':
#         return jsonify({'error': 'No job description file selected'}), 400
    
#     print(f"Processing files: {resume_file.filename}, {job_description_file.filename}")
    
#     resume_path = None
#     job_path = None
    
#     try:
#         # Secure filenames and save files
#         resume_filename = secure_filename(resume_file.filename)
#         job_filename = secure_filename(job_description_file.filename)
        
#         resume_path = os.path.join(app.config['UPLOAD_FOLDER'], resume_filename)
#         job_path = os.path.join(app.config['UPLOAD_FOLDER'], job_filename)
        
#         resume_file.save(resume_path)
#         job_description_file.save(job_path)
        
#         print("Files saved successfully")
        
#         # Extract text from files
#         resume_text = extract_text(resume_path)
#         job_text = extract_text(job_path)
        
#         print(f"Resume text length: {len(resume_text)}")
#         print(f"Job description text length: {len(job_text)}")
        
#         if len(resume_text.strip()) == 0:
#             return jsonify({'error': 'Could not extract text from resume file'}), 400
#         if len(job_text.strip()) == 0:
#             return jsonify({'error': 'Could not extract text from job description file'}), 400
        
#         # Analyze skills
#         skills_analysis = compare_skills(resume_text, job_text)
        
#         print("Skill analysis completed successfully")
#         return jsonify(skills_analysis)
        
#     except Exception as e:
#         print(f"Error in skill analysis: {str(e)}")
#         return jsonify({'error': f'An error occurred during skill analysis: {str(e)}'}), 500
#     finally:
#         # Clean up files
#         if resume_path and os.path.exists(resume_path):
#             os.remove(resume_path)
#             print("Resume file cleaned up")
#         if job_path and os.path.exists(job_path):
#             os.remove(job_path)
#             print("Job description file cleaned up")

# def extract_text(file_path):
#     print(f"Extracting text from: {file_path}")
#     _, file_extension = os.path.splitext(file_path)
    
#     try:
#         if file_extension.lower() == '.pdf':
#             with open(file_path, 'rb') as file:
#                 reader = PyPDF2.PdfReader(file)
#                 text = ''
#                 for page in reader.pages:
#                     page_text = page.extract_text()
#                     if page_text:
#                         text += page_text + ' '
#                 print(f"PDF extraction completed: {len(text)} characters")
                
#         elif file_extension.lower() in ['.docx', '.doc']:
#             doc = docx.Document(file_path)
#             text = '\n'.join([paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()])
#             print(f"DOCX extraction completed: {len(text)} characters")
            
#         else:
#             # Try to read as text file
#             try:
#                 with open(file_path, 'r', encoding='utf-8') as file:
#                     text = file.read()
#             except:
#                 with open(file_path, 'r', encoding='latin-1') as file:
#                     text = file.read()
#             print(f"Text file extraction completed: {len(text)} characters")
        
#         return text.strip()
        
#     except Exception as e:
#         print(f"Error extracting text from {file_path}: {str(e)}")
#         return ""

# def compare_skills(resume_text, job_text):
#     print("Comparing skills using Groq API...")
    
#     # If Groq client is not available, return mock data
#     if not groq_client:
#         print("Groq client not available, returning mock data")
#         return {
#             "skills_from_resume": ["Python", "JavaScript", "React", "Node.js", "MongoDB", "HTML5", "CSS3", "Git"],
#             "skills_required_in_job": ["Python", "React", "AWS", "Docker", "Kubernetes", "HTML5", "CSS3", "JavaScript"],
#             "matching_skills": ["Python", "React", "JavaScript", "HTML5", "CSS3"],
#             "skills_to_improve": ["AWS", "Docker", "Kubernetes"]
#         }
    
#     prompt = f"""
#     Analyze the following resume and job description to extract and compare skills.

#     RESUME TEXT:
#     {resume_text[:3000]}

#     JOB DESCRIPTION TEXT:
#     {job_text[:3000]}

#     Please provide a JSON response with exactly these 4 arrays:
#     1. "skills_from_resume" - technical skills found in the resume
#     2. "skills_required_in_job" - specific technical skills required in the job description
#     3. "matching_skills" - skills that appear in both lists
#     4. "skills_to_improve" - skills from job description not found in resume

#     Return ONLY valid JSON, no other text.
#     Example format:
#     {{
#         "skills_from_resume": ["Python", "JavaScript", "React"],
#         "skills_required_in_job": ["Python", "React", "AWS"],
#         "matching_skills": ["Python", "React"],
#         "skills_to_improve": ["AWS"]
#     }}
#     """

#     try:
#         print("Sending request to Groq API...")
#         chat_completion = groq_client.chat.completions.create(
#             messages=[
#                 {
#                     "role": "user",
#                     "content": prompt
#                 }
#             ],
#             # model="llama3-70b-8192",
#             model="openai/gpt-oss-120b",
#             temperature=0.1,
#             max_tokens=1000
#         )
        
#         response = chat_completion.choices[0].message.content
#         print("Raw API response received")
        
#         # Clean the response
#         response = response.strip()
        
#         # Remove markdown code blocks if present
#         if response.startswith('```json'):
#             response = response[7:]
#         if response.startswith('```'):
#             response = response[3:]
#         if response.endswith('```'):
#             response = response[:-3]
        
#         response = response.strip()
#         print(f"Cleaned response: {response[:200]}...")
        
#         # Parse JSON
#         skills_data = json.loads(response)
        
#         # Validate required keys
#         required_keys = ["skills_from_resume", "skills_required_in_job", "matching_skills", "skills_to_improve"]
#         for key in required_keys:
#             if key not in skills_data:
#                 skills_data[key] = []
        
#         print("Skill analysis successful")
#         return skills_data

#     except json.JSONDecodeError as e:
#         print(f"JSON Decode Error: {e}")
#         print(f"Problematic response: {response}")
#         # Return fallback data
#         return {
#             "skills_from_resume": ["Python", "JavaScript", "React", "HTML5", "CSS3"],
#             "skills_required_in_job": ["Python", "React", "AWS", "Docker", "HTML5", "CSS3"],
#             "matching_skills": ["Python", "React", "HTML5", "CSS3"],
#             "skills_to_improve": ["AWS", "Docker"]
#         }
#     except Exception as e:
#         print(f"Error in compare_skills: {e}")
#         # Return fallback data
#         return {
#             "skills_from_resume": ["Python", "JavaScript", "React", "HTML5", "CSS3"],
#             "skills_required_in_job": ["Python", "React", "AWS", "Docker", "HTML5", "CSS3"],
#             "matching_skills": ["Python", "React", "HTML5", "CSS3"],
#             "skills_to_improve": ["AWS", "Docker"]
#         }

# # ==================== COURSE RECOMMENDATION ====================

# @app.route('/recommend_course', methods=['POST'])
# def recommend_course_api():
#     try:
#         print("=== RECOMMEND COURSE ENDPOINT CALLED ===")
#         data = request.get_json()
        
#         if not data:
#             return jsonify({'error': 'No JSON data received'}), 400
            
#         skill_name = data.get('resource')
        
#         if not skill_name:
#             return jsonify({'error': 'Skill name is required'}), 400
        
#         print(f"Course recommendation requested for: {skill_name}")
        
#         recommended_link = recommend_course(skill_name)
        
#         # Handle pandas Series
#         if isinstance(recommended_link, pd.Series):
#             if recommended_link.empty:
#                 return jsonify({'error': 'No recommendation found'}), 404
#             recommended_link = recommended_link.iloc[0]
#         elif not recommended_link:
#             return jsonify({'error': 'No recommendation found'}), 404
        
#         # Ensure it's a string
#         recommended_link = str(recommended_link).strip()
        
#         # Validate URL format
#         if not recommended_link.startswith(('http://', 'https://')):
#             if recommended_link.startswith('www.'):
#                 recommended_link = 'https://' + recommended_link
#             else:
#                 recommended_link = f'https://www.udemy.com{recommended_link if recommended_link.startswith("/") else "/" + recommended_link}'
        
#         print(f"Final recommended link: {recommended_link}")
#         return jsonify({'recommendation': recommended_link})
    
#     except Exception as e:
#         print(f"Error in course recommendation: {e}")
#         # Return a fallback URL
#         skill = data.get('resource', 'programming') if 'data' in locals() else 'programming'
#         fallback_url = f'https://www.udemy.com/courses/search/?src=ukw&q={skill.replace(" ", "+")}'
#         return jsonify({'recommendation': fallback_url})

# # ==================== CHATBOT ROUTE ====================

# @app.route('/api/chat', methods=['POST'])
# def chat_with_bot():
#     try:
#         print("=== CHAT ENDPOINT CALLED ===")
#         data = request.get_json()
#         user_message = data.get('message', '').strip().lower()
        
#         print(f"User message: {user_message}")
        
#         # Simple rule-based responses
#         if not user_message:
#             response = "Hello! I'm SkillSpark AI Assistant. How can I help you today?"
#         elif any(word in user_message for word in ['hello', 'hi', 'hey']):
#             response = "Hello! 👋 I'm SkillSpark AI Assistant. I can help you with skill analysis, career guidance, and learning recommendations!"
#         elif any(word in user_message for word in ['skill', 'learn', 'study']):
#             response = "I recommend focusing on: Python, JavaScript, React, Cloud technologies, and AI/ML skills! 🚀"
#         elif any(word in user_message for word in ['resume', 'cv']):
#             response = "For resume tips: Highlight projects, quantify achievements, and tailor to job descriptions! 📄"
#         elif any(word in user_message for word in ['job', 'career', 'interview']):
#             response = "Job search advice: Network actively, practice interviews, and build a strong portfolio! 💼"
#         elif any(word in user_message for word in ['thank']):
#             response = "You're welcome! 😊 Happy to help with your career journey!"
#         else:
#             response = "I can help with: Skills development, Resume improvement, Job search strategies, and Career guidance! What would you like to know? 🎯"
        
#         print(f"Bot response: {response}")
        
#         return jsonify({
#             'response': response,
#             'timestamp': datetime.now().isoformat()
#         })
        
#     except Exception as e:
#         print(f"Chat error: {e}")
#         return jsonify({
#             'response': "Hello! I'm here to help with your career development. Ask me anything!",
#             'timestamp': datetime.now().isoformat()
#         })

# # ==================== TEST ROUTES ====================

# @app.route('/api/health', methods=['GET'])
# def health_check():
#     return jsonify({
#         'status': 'healthy',
#         'mongodb': 'connected' if client else 'disconnected',
#         'groq': 'available' if groq_client else 'unavailable',
#         'timestamp': datetime.now().isoformat()
#     })

# @app.route('/api/test', methods=['GET'])
# def test_endpoint():
#     return jsonify({'message': 'Backend is working!'})

# @app.route('/api/test_chat', methods=['GET'])
# def test_chat():
#     return jsonify({'message': 'Chatbot endpoint is working!'})

# # ==================== MAIN ====================

# if __name__ == '__main__':
#     print("=== STARTING FLASK SERVER ===")
#     print(f"Upload folder: {os.path.abspath(UPLOAD_FOLDER)}")
#     print(f"MongoDB: {'Connected' if client else 'Not connected'}")
#     print(f"Groq API: {'Available' if groq_client else 'Not available'}")
#     app.run(host='0.0.0.0', port=5000, debug=True)










from flask import Flask, jsonify, request
from flask_cors import CORS
from pymongo import MongoClient
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from datetime import datetime
import os
import PyPDF2
import docx
import json
import re
from groq import Groq
from google import genai
import pandas as pd
from rec_courses import recommend_course
from dotenv import load_dotenv


# ==================== APP INITIALIZATION ====================

app = Flask(__name__)
CORS(app)
load_dotenv()


# ==================== ENVIRONMENT VARIABLES ====================

connec_string = os.getenv('connec_string')
groq_api = os.getenv('GROQ_API_KEY')
gemini_api = os.getenv('GEMINI_API_KEY')

print(f"Connection string exists: {bool(connec_string)}")
print(f"Groq API key exists: {bool(groq_api)}")
print(f"Gemini API key exists: {bool(gemini_api)}")


# ==================== UPLOAD CONFIGURATION ====================

UPLOAD_FOLDER = 'uploads'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)
    print(f"Created uploads directory: {UPLOAD_FOLDER}")


# ==================== MONGODB CONFIGURATION ====================

client = None
users_collection = None

if connec_string:
    try:
        client = MongoClient(connec_string)

        # Test MongoDB connection
        client.admin.command('ping')

        print("MongoDB is connected")

        db = client['UserTest']
        users_collection = db['users']

        print(f"Database connected: {db.name}")

    except Exception as e:
        print(f"MongoDB connection failed: {e}")
        client = None
        users_collection = None

else:
    print("No MongoDB connection string found")
    users_collection = None


# ==================== GROQ CONFIGURATION ====================

groq_client = None

if groq_api:
    try:
        groq_client = Groq(api_key=groq_api)
        print("Groq client initialized successfully")

    except Exception as e:
        print(f"Groq client initialization failed: {e}")
        groq_client = None

else:
    print("No Groq API key found")
    groq_client = None


# ==================== GEMINI CONFIGURATION ====================

gemini_client = None

if gemini_api:
    try:
        gemini_client = genai.Client(api_key=gemini_api)
        print("Gemini client initialized successfully")

    except Exception as e:
        print(f"Gemini client initialization failed: {e}")
        gemini_client = None

else:
    print("No Gemini API key found")
    gemini_client = None


# ==================== AUTHENTICATION ROUTES ====================


@app.route('/api/signup', methods=['POST'])
def signup():

    print("=== SIGNUP ENDPOINT CALLED ===")

    try:

        if not request.is_json:
            return jsonify({
                'message': 'Request must be JSON'
            }), 400

        data = request.get_json()

        print("Signup request received")

        name = data.get('name')
        email = data.get('email')
        password = data.get('password')

        if not name or not email or not password:
            return jsonify({
                'message': 'Name, email, and password are required'
            }), 400

        # Check if user already exists
        if users_collection is not None:

            existing_user = users_collection.find_one({
                'email': email
            })

            if existing_user:
                return jsonify({
                    'message': 'User already exists with this email'
                }), 400

            # Hash password
            hashed_password = generate_password_hash(password)

            user_data = {
                'name': name,
                'email': email,
                'password': hashed_password,
                'created_at': datetime.now().isoformat(),
                'skills': [],
                'job': '',
                'tagline': 'A catchy tagline!'
            }

            result = users_collection.insert_one(user_data)

            print(f"User created with ID: {result.inserted_id}")

            return jsonify({
                'message': 'Signup successful',
                'user': {
                    'name': name,
                    'email': email
                }
            }), 201

        else:

            return jsonify({
                'message': 'Database not available'
            }), 500

    except Exception as e:

        print(f"Signup error: {e}")

        return jsonify({
            'message': 'Error creating user'
        }), 500


# ==================== LOGIN ====================


@app.route('/api/login', methods=['POST'])
def login():

    print("=== LOGIN ENDPOINT CALLED ===")

    try:

        if not request.is_json:
            return jsonify({
                'message': 'Request must be JSON'
            }), 400

        data = request.get_json()

        print("Login request received")

        email = data.get('email')
        password = data.get('password')

        if not email or not password:
            return jsonify({
                'message': 'Email and password are required'
            }), 400

        if users_collection is not None:

            user = users_collection.find_one({
                'email': email
            })

            if user and check_password_hash(
                user['password'],
                password
            ):

                user_data = {
                    'name': user['name'],
                    'email': user['email'],
                    'id': str(user['_id'])
                }

                return jsonify({
                    'message': 'Login successful',
                    'user': user_data
                }), 200

            else:

                return jsonify({
                    'message': 'Invalid email or password'
                }), 401

        else:

            return jsonify({
                'message': 'Database not available'
            }), 500

    except Exception as e:

        print(f"Login error: {e}")

        return jsonify({
            'message': 'Error during login'
        }), 500


# ==================== GET USER PROFILE ====================


@app.route('/api/user/profile', methods=['GET'])
def get_user_profile():

    try:

        email = request.args.get('email')

        if not email:
            return jsonify({
                'message': 'Email is required'
            }), 400

        if users_collection is not None:

            user = users_collection.find_one({
                'email': email
            })

            if user:

                user_data = {
                    'name': user.get('name', ''),
                    'email': user.get('email', ''),
                    'skills': user.get('skills', []),
                    'job': user.get('job', ''),
                    'tagline': user.get(
                        'tagline',
                        'A catchy tagline!'
                    )
                }

                return jsonify({
                    'user': user_data
                }), 200

            else:

                return jsonify({
                    'message': 'User not found'
                }), 404

        else:

            return jsonify({
                'message': 'Database not available'
            }), 500

    except Exception as e:

        print(f"Profile error: {e}")

        return jsonify({
            'message': 'Error fetching profile'
        }), 500


# ==================== UPDATE USER PROFILE ====================


@app.route('/api/user/profile', methods=['PUT'])
def update_user_profile():

    try:

        data = request.get_json()

        if not data:
            return jsonify({
                'message': 'No data received'
            }), 400

        email = data.get('email')
        updates = data.get('updates', {})

        if not email or not updates:
            return jsonify({
                'message': 'Email and updates are required'
            }), 400

        if users_collection is not None:

            # Never allow password update through profile route
            updates.pop('password', None)

            result = users_collection.update_one(
                {'email': email},
                {'$set': updates}
            )

            if result.modified_count:

                return jsonify({
                    'message': 'Profile updated successfully'
                }), 200

            else:

                return jsonify({
                    'message': 'No changes made or user not found'
                }), 400

        else:

            return jsonify({
                'message': 'Database not available'
            }), 500

    except Exception as e:

        print(f"Update profile error: {e}")

        return jsonify({
            'message': 'Error updating profile'
        }), 500


# ==================== SKILL ANALYZER ROUTE ====================


@app.route('/api/skill-analyzer', methods=['POST'])
def skill_analyzer():

    print("=== SKILL ANALYZER ENDPOINT CALLED ===")

    # Check resume
    if 'resume' not in request.files:

        print("Missing resume file")

        return jsonify({
            'error': 'Resume file is required'
        }), 400

    # Check job description
    if 'job_description' not in request.files:

        print("Missing job description file")

        return jsonify({
            'error': 'Job description file is required'
        }), 400

    resume_file = request.files['resume']
    job_description_file = request.files['job_description']

    # Check filenames
    if resume_file.filename == '':

        return jsonify({
            'error': 'No resume file selected'
        }), 400

    if job_description_file.filename == '':

        return jsonify({
            'error': 'No job description file selected'
        }), 400

    print(
        f"Processing files: "
        f"{resume_file.filename}, "
        f"{job_description_file.filename}"
    )

    resume_path = None
    job_path = None

    try:

        # Secure filenames
        resume_filename = secure_filename(
            resume_file.filename
        )

        job_filename = secure_filename(
            job_description_file.filename
        )

        resume_path = os.path.join(
            app.config['UPLOAD_FOLDER'],
            resume_filename
        )

        job_path = os.path.join(
            app.config['UPLOAD_FOLDER'],
            job_filename
        )

        # Save files
        resume_file.save(resume_path)
        job_description_file.save(job_path)

        print("Files saved successfully")

        # Extract text
        resume_text = extract_text(resume_path)
        job_text = extract_text(job_path)

        print(
            f"Resume text length: "
            f"{len(resume_text)}"
        )

        print(
            f"Job description text length: "
            f"{len(job_text)}"
        )

        if len(resume_text.strip()) == 0:

            return jsonify({
                'error':
                'Could not extract text from resume file'
            }), 400

        if len(job_text.strip()) == 0:

            return jsonify({
                'error':
                'Could not extract text from job description file'
            }), 400

        # Analyze skills using Groq
        skills_analysis = compare_skills(
            resume_text,
            job_text
        )

        print(
            "Skill analysis completed successfully"
        )

        return jsonify(skills_analysis)

    except Exception as e:

        print(
            f"Error in skill analysis: {str(e)}"
        )

        return jsonify({
            'error':
            f'An error occurred during skill analysis: {str(e)}'
        }), 500

    finally:

        # Remove resume
        if resume_path and os.path.exists(resume_path):

            os.remove(resume_path)

            print("Resume file cleaned up")

        # Remove job description
        if job_path and os.path.exists(job_path):

            os.remove(job_path)

            print("Job description file cleaned up")


# ==================== TEXT EXTRACTION ====================


def extract_text(file_path):

    print(
        f"Extracting text from: {file_path}"
    )

    _, file_extension = os.path.splitext(
        file_path
    )

    try:

        # PDF
        if file_extension.lower() == '.pdf':

            with open(
                file_path,
                'rb'
            ) as file:

                reader = PyPDF2.PdfReader(file)

                text = ''

                for page in reader.pages:

                    page_text = page.extract_text()

                    if page_text:

                        text += page_text + ' '

                print(
                    f"PDF extraction completed: "
                    f"{len(text)} characters"
                )

        # DOCX / DOC
        elif file_extension.lower() in [
            '.docx',
            '.doc'
        ]:

            doc = docx.Document(file_path)

            text = '\n'.join(
                [
                    paragraph.text
                    for paragraph in doc.paragraphs
                    if paragraph.text.strip()
                ]
            )

            print(
                f"DOCX extraction completed: "
                f"{len(text)} characters"
            )

        # TXT / Other
        else:

            try:

                with open(
                    file_path,
                    'r',
                    encoding='utf-8'
                ) as file:

                    text = file.read()

            except:

                with open(
                    file_path,
                    'r',
                    encoding='latin-1'
                ) as file:

                    text = file.read()

            print(
                f"Text file extraction completed: "
                f"{len(text)} characters"
            )

        return text.strip()

    except Exception as e:

        print(
            f"Error extracting text from "
            f"{file_path}: {str(e)}"
        )

        return ""


# ==================== GROQ SKILL COMPARISON ====================


def compare_skills(
    resume_text,
    job_text
):

    print(
        "Comparing skills using Groq API..."
    )

    # If Groq unavailable
    if not groq_client:

        print(
            "Groq client not available, "
            "returning mock data"
        )

        return {

            "skills_from_resume": [
                "Python",
                "JavaScript",
                "React",
                "Node.js",
                "MongoDB",
                "HTML5",
                "CSS3",
                "Git"
            ],

            "skills_required_in_job": [
                "Python",
                "React",
                "AWS",
                "Docker",
                "Kubernetes",
                "HTML5",
                "CSS3",
                "JavaScript"
            ],

            "matching_skills": [
                "Python",
                "React",
                "JavaScript",
                "HTML5",
                "CSS3"
            ],

            "skills_to_improve": [
                "AWS",
                "Docker",
                "Kubernetes"
            ]
        }

    prompt = f"""
Analyze the following resume and job description
to extract and compare skills.

RESUME TEXT:
{resume_text[:3000]}

JOB DESCRIPTION TEXT:
{job_text[:3000]}

Please provide a JSON response with exactly
these 4 arrays:

1. "skills_from_resume"
- technical skills found in the resume

2. "skills_required_in_job"
- specific technical skills required in the job description

3. "matching_skills"
- skills that appear in both lists

4. "skills_to_improve"
- skills from job description not found in resume

Return ONLY valid JSON.
Do not include markdown.
Do not include explanations.

Example:

{{
    "skills_from_resume": [
        "Python",
        "JavaScript",
        "React"
    ],

    "skills_required_in_job": [
        "Python",
        "React",
        "AWS"
    ],

    "matching_skills": [
        "Python",
        "React"
    ],

    "skills_to_improve": [
        "AWS"
    ]
}}
"""

    try:

        print(
            "Sending request to Groq API..."
        )

        chat_completion = (
            groq_client.chat.completions.create(

                messages=[
                    {
                        "role": "user",
                        "content": prompt
                    }
                ],

                model="openai/gpt-oss-120b",

                temperature=0.1,

                max_tokens=1000
            )
        )

        response = (
            chat_completion
            .choices[0]
            .message
            .content
        )

        print(
            "Raw API response received"
        )

        # Clean response
        response = response.strip()

        # Remove markdown JSON block
        if response.startswith(
            '```json'
        ):

            response = response[7:]

        if response.startswith(
            '```'
        ):

            response = response[3:]

        if response.endswith(
            '```'
        ):

            response = response[:-3]

        response = response.strip()

        print(
            f"Cleaned response: "
            f"{response[:200]}..."
        )

        # Parse JSON
        skills_data = json.loads(response)

        # Validate keys
        required_keys = [
            "skills_from_resume",
            "skills_required_in_job",
            "matching_skills",
            "skills_to_improve"
        ]

        for key in required_keys:

            if key not in skills_data:

                skills_data[key] = []

        print(
            "Skill analysis successful"
        )

        return skills_data

    except json.JSONDecodeError as e:

        print(
            f"JSON Decode Error: {e}"
        )

        print(
            f"Problematic response: {response}"
        )

        return {

            "skills_from_resume": [
                "Python",
                "JavaScript",
                "React",
                "HTML5",
                "CSS3"
            ],

            "skills_required_in_job": [
                "Python",
                "React",
                "AWS",
                "Docker",
                "HTML5",
                "CSS3"
            ],

            "matching_skills": [
                "Python",
                "React",
                "HTML5",
                "CSS3"
            ],

            "skills_to_improve": [
                "AWS",
                "Docker"
            ]
        }

    except Exception as e:

        print(
            f"Error in compare_skills: {e}"
        )

        return {

            "skills_from_resume": [
                "Python",
                "JavaScript",
                "React",
                "HTML5",
                "CSS3"
            ],

            "skills_required_in_job": [
                "Python",
                "React",
                "AWS",
                "Docker",
                "HTML5",
                "CSS3"
            ],

            "matching_skills": [
                "Python",
                "React",
                "HTML5",
                "CSS3"
            ],

            "skills_to_improve": [
                "AWS",
                "Docker"
            ]
        }


# ==================== COURSE RECOMMENDATION ====================


@app.route('/recommend_course', methods=['POST'])
def recommend_course_api():

    try:

        print(
            "=== RECOMMEND COURSE ENDPOINT CALLED ==="
        )

        data = request.get_json()

        if not data:

            return jsonify({
                'error': 'No JSON data received'
            }), 400

        skill_name = data.get('resource')

        if not skill_name:

            return jsonify({
                'error': 'Skill name is required'
            }), 400

        print(
            f"Course recommendation requested "
            f"for: {skill_name}"
        )

        recommended_link = recommend_course(
            skill_name
        )

        # Handle pandas Series
        if isinstance(
            recommended_link,
            pd.Series
        ):

            if recommended_link.empty:

                return jsonify({
                    'error': 'No recommendation found'
                }), 404

            recommended_link = (
                recommended_link.iloc[0]
            )

        elif not recommended_link:

            return jsonify({
                'error': 'No recommendation found'
            }), 404

        # Convert to string
        recommended_link = str(
            recommended_link
        ).strip()

        # Validate URL
        if not recommended_link.startswith(
            ('http://', 'https://')
        ):

            if recommended_link.startswith(
                'www.'
            ):

                recommended_link = (
                    'https://' +
                    recommended_link
                )

            else:

                recommended_link = (
                    'https://www.udemy.com'
                    +
                    (
                        recommended_link
                        if recommended_link.startswith('/')
                        else '/' + recommended_link
                    )
                )

        print(
            f"Final recommended link: "
            f"{recommended_link}"
        )

        return jsonify({
            'recommendation':
            recommended_link
        })

    except Exception as e:

        print(
            f"Error in course recommendation: {e}"
        )

        skill = (
            data.get('resource', 'programming')
            if 'data' in locals()
            else 'programming'
        )

        fallback_url = (
            'https://www.udemy.com/courses/search/'
            '?src=ukw&q='
            +
            skill.replace(" ", "+")
        )

        return jsonify({
            'recommendation':
            fallback_url
        })


# ==================== GEMINI CHATBOT ROUTE ====================

@app.route('/api/chat', methods=['POST'])
def chat_with_bot():

    try:

        print("=== CHAT ENDPOINT CALLED ===")

        data = request.get_json()

        if not data:
            return jsonify({
                'response': 'Please send a message.',
                'timestamp': datetime.now().isoformat()
            }), 400

        user_message = data.get('message', '').strip()
        history = data.get('history', [])

        print(f"User message: {user_message}")

        if not user_message:
            return jsonify({
                'response': 'Please type a message.',
                'timestamp': datetime.now().isoformat()
            }), 400

        # Check Gemini
        if not gemini_client:
            print("Gemini client not available")

            return jsonify({
                'response':
                    'Gemini AI is currently unavailable. '
                    'Please try again later.',
                'timestamp': datetime.now().isoformat()
            }), 503

        # ==================== BUILD CHAT HISTORY ====================

        conversation = []

        for message in history[-10:]:

            role = message.get('type')

            content = message.get(
                'content',
                ''
            ).strip()

            if not content:
                continue

            if role == 'user':
                conversation.append(
                    f"User: {content}"
                )

            elif role == 'bot':
                conversation.append(
                    f"Assistant: {content}"
                )

        # Add current message
        conversation.append(
            f"User: {user_message}"
        )

        conversation_text = "\n".join(
            conversation
        )

        # ==================== GEMINI PROMPT ====================

        prompt = f"""
You are SkillSpark AI Assistant,
an AI career and learning assistant.

Help students and job seekers with:

- Skill gap analysis
- Programming
- Technical skills
- Resume and CV improvement
- Interview preparation
- Career guidance
- Learning roadmaps
- Course recommendations
- Software development
- Artificial Intelligence
- Machine Learning
- Web development
- Cloud computing
- DevOps
- Python
- JavaScript
- React
- Node.js
- MongoDB
- Data Structures and Algorithms

Give clear, practical and beginner-friendly answers.

For technical questions:
- Explain concepts simply
- Give examples when useful
- Give step-by-step guidance

For career questions:
- Give structured and practical advice
- Mention useful skills and technologies

Do not claim that you analyzed a resume
or job description unless that information
was actually provided.

Be friendly and professional.

Conversation history:

{conversation_text}

Respond naturally to the user's latest message.
"""

        # ==================== GEMINI REQUEST WITH RETRY ====================

        max_attempts = 3

        for attempt in range(1, max_attempts + 1):

            try:

                print(
                    f"Sending request to Gemini API "
                    f"(attempt {attempt}/{max_attempts})..."
                )

                response = gemini_client.models.generate_content(
                    model="gemini-3.6-flash",
                    contents=prompt
                )

                bot_response = (
                    response.text.strip()
                    if response.text
                    else
                    "Sorry, I couldn't generate a response."
                )

                print(
                    "Gemini response received successfully"
                )

                return jsonify({
                    'response': bot_response,
                    'timestamp': datetime.now().isoformat()
                }), 200

            except Exception as gemini_error:

                error_text = str(gemini_error)

                print(
                    f"Gemini attempt {attempt} failed:"
                )

                print(error_text)

                # Retry only for temporary server problems
                if (
                    '503' in error_text
                    or
                    'UNAVAILABLE' in error_text
                    or
                    '429' in error_text
                ):

                    if attempt < max_attempts:

                        import time

                        wait_time = 2 ** attempt

                        print(
                            f"Temporary Gemini error. "
                            f"Retrying in {wait_time} seconds..."
                        )

                        time.sleep(wait_time)

                        continue

                    else:

                        print(
                            "Gemini unavailable after "
                            "all retry attempts."
                        )

                        return jsonify({
                            'response':
                                'Gemini AI is temporarily '
                                'busy right now. '
                                'Please try your message '
                                'again in a few seconds.',
                            'timestamp':
                                datetime.now().isoformat()
                        }), 503

                # Don't retry authentication/configuration errors
                else:

                    print(
                        "Non-retryable Gemini error."
                    )

                    return jsonify({
                        'response':
                            'There was a problem connecting '
                            'to the AI service. '
                            'Please try again later.',
                        'timestamp':
                            datetime.now().isoformat()
                    }), 500

    except Exception as e:

        print(
            f"Chat endpoint error: {e}"
        )

        return jsonify({
            'response':
                "Sorry, I couldn't process your "
                "request right now. Please try again.",
            'timestamp':
                datetime.now().isoformat()
        }), 500

# ==================== HEALTH CHECK ====================


@app.route('/api/health', methods=['GET'])
def health_check():

    return jsonify({

        'status':
        'healthy',

        'mongodb':
        'connected'
        if client
        else
        'disconnected',

        'groq':
        'available'
        if groq_client
        else
        'unavailable',

        'gemini':
        'available'
        if gemini_client
        else
        'unavailable',

        'timestamp':
        datetime.now().isoformat()
    })


# ==================== TEST ROUTES ====================


@app.route('/api/test', methods=['GET'])
def test_endpoint():

    return jsonify({
        'message':
        'Backend is working!'
    })


@app.route('/api/test_chat', methods=['GET'])
def test_chat():

    return jsonify({
        'message':
        'Chatbot endpoint is working!'
    })


# ==================== MAIN ====================


if __name__ == '__main__':

    print(
        "=== STARTING FLASK SERVER ==="
    )

    print(
        f"Upload folder: "
        f"{os.path.abspath(UPLOAD_FOLDER)}"
    )

    print(
        f"MongoDB: "
        f"{'Connected' if client else 'Not connected'}"
    )

    print(
        f"Groq API: "
        f"{'Available' if groq_client else 'Not available'}"
    )

    print(
        f"Gemini API: "
        f"{'Available' if gemini_client else 'Not available'}"
    )

    app.run(
        host='0.0.0.0',
        port=5000,
        debug=True
    )